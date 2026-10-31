import os
import sys
import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

# 本脚本位于 <项目>/03-EDA导出/工具/，默认读取上一级目录中的 BOM
EDA_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT = os.path.join(EDA_DIR, "BOM_Board1_PCB1_2026-10-11.xlsx")
p = sys.argv[1] if len(sys.argv) > 1 else DEFAULT
wb = openpyxl.load_workbook(p, data_only=False)
for ws in wb.worksheets:
    print("=== SHEET:", ws.title, "| dims:", ws.dimensions, "| rows:", ws.max_row, "| cols:", ws.max_column)
    for r in ws.iter_rows(min_row=1, max_row=ws.max_row):
        vals = []
        for c in r:
            v = c.value
            vals.append("" if v is None else str(v).replace("\n", " "))
        line = " | ".join(vals).rstrip(" |")
        if line.strip():
            print(f"{r[0].row:>3}: {line}")
    print()
